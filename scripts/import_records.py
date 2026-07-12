#!/usr/bin/env python3
"""Merge IABA's records workbook (CW.xlsx) into data/plots.geojson.

Re-runnable: statuses/colors are synced from data/reference/layout-spaces.json
(the digitized diagram, incl. COLOR_OVERRIDES for post-PDF changes), spaces
missing from the map are appended using the placement transform stored in
plots.geojson metadata, then ownership/burial data is merged from the workbook:

  - vacant plots      -> properties.owner ("LAST, First" of the deed holder)
  - burials           -> properties.person = {name, burial: ISO date}

Notes columns (Deeded/AB/anecdotes) are intentionally NOT exported — internal.
Owner names are only published on vacant plots (owner approved 2026-07-12);
owners of used spaces stay out of the public file.

Usage: python3 scripts/import_records.py [path/to/CW.xlsx]
"""

import json
import math
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
PLOTS = ROOT / "data/plots.geojson"
LAYOUT = ROOT / "data/reference/layout-spaces.json"
DEFAULT_XLSX = Path.home() / "Downloads/CW.xlsx"

# Owner-confirmed corrections (2026-07-12) where the workbook's All/Burials
# sheets disagree or contain typos. Keyed by workbook space ID.
CORRECTIONS: dict[str, dict[str, str]] = {
    "48A-1": {"owner": "SHARAFUDDIN, Ossamah", "deceased": "SHARAFUDDIN, Essam"},
    "28C-3": {"owner": "BAHRAMI, Behzad", "deceased": "BAHRAMI, Sultan Ali"},
    "6C-3": {"deceased": "SHAFIGHI, Zeenat"},
    "6C-4": {"deceased": "SHAFIGHI, Mohamed"},
    "29B-2": {"deceased": "DAWOODALY, Mohamedali"},
    "18C-2": {"owner": "MOADDEB, Mahshid"},
}

M_PER_DEG_LAT = 111_320


def to_plot_id(sheet_id: str) -> str:
    """Workbook '4A-2' -> map '04-A2'."""
    m = re.fullmatch(r"(\d+)([A-D])-([1-4])", sheet_id.strip())
    if not m:
        raise ValueError(f"unparseable space ID: {sheet_id!r}")
    return f"{int(m.group(1)):02d}-{m.group(2)}{m.group(3)}"


def place_rect(rect, origin, params):
    """Replicate src/layout.ts toLatLng for one rect_px -> polygon ring."""
    cx, cy = origin
    anchor = params["anchor"]
    sx = params["spaceWidthM"] / 70.0
    sy = params["spaceDepthM"] / 103.0
    theta = math.radians(params["rotationDeg"])
    cos, sin = math.cos(theta), math.sin(theta)
    m_per_deg_lng = M_PER_DEG_LAT * math.cos(math.radians(anchor["lat"]))

    def ll(px, py):
        e0 = (px - cx) * sx
        n0 = -(py - cy) * sy
        e = e0 * cos + n0 * sin
        n = -e0 * sin + n0 * cos
        return [anchor["lng"] + e / m_per_deg_lng, anchor["lat"] + n / M_PER_DEG_LAT]

    x0, y0, x1, y1 = rect
    ring = [ll(x0, y0), ll(x1, y0), ll(x1, y1), ll(x0, y1)]
    return ring + [ring[0]]


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    fc = json.loads(PLOTS.read_text())
    layout = json.loads(LAYOUT.read_text())
    params = fc["metadata"]["params"]
    plots = {f["properties"]["id"]: f for f in fc["features"]}

    # -- 1. sync statuses/colors from the digitized layout; append new spaces --
    xs0, ys0, xs1, ys1 = zip(*(s["rect_px"] for s in layout["spaces"]))
    origin = ((min(xs0) + max(xs1)) / 2, (min(ys0) + max(ys1)) / 2)
    changed, added = [], []
    for s in layout["spaces"]:
        f = plots.get(s["id"])
        if f is None:
            plots[s["id"]] = f = {
                "type": "Feature",
                "properties": {
                    "id": s["id"], "lot": s["lot"], "space": s["space"],
                    "status": s["status"], "color": s["color"],
                },
                "geometry": {"type": "Polygon",
                             "coordinates": [place_rect(s["rect_px"], origin, params)]},
            }
            fc["features"].append(f)
            added.append(s["id"])
        elif f["properties"]["status"] != s["status"]:
            changed.append(f"{s['id']}: {f['properties']['status']} -> {s['status']}")
            f["properties"]["status"] = s["status"]
            f["properties"]["color"] = s["color"]

    # -- 2. merge ownership/burials from the workbook --
    wb = openpyxl.load_workbook(xlsx, data_only=True)

    def clean(v):
        v = str(v).strip() if v is not None else ""
        return v or None

    records = {}  # plot id -> {owner, person?}
    for r in wb["All"].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[2], str) or not clean(r[2]):
            continue
        sid = r[2].strip()
        fix = CORRECTIONS.get(sid, {})
        owner = fix.get("owner") or clean(r[5])
        used = r[3] == 1
        rec = {"owner": owner, "used": used}
        if used:
            name = fix.get("deceased") or clean(r[6])
            if not name:
                sys.exit(f"ERROR: {sid} is used but has no deceased name")
            rec["person"] = {"name": name, "burial": r[4].date().isoformat()}
        records[to_plot_id(sid)] = rec

    # burials missing from the All sheet (e.g. 59D-4 in the future-lot area)
    for r in wb["Burials"].iter_rows(min_row=3, values_only=True):
        if not isinstance(r[2], str) or not clean(r[2]):
            continue
        pid = to_plot_id(r[2])
        if pid not in records:
            fix = CORRECTIONS.get(r[2].strip(), {})
            records[pid] = {
                "owner": fix.get("owner") or clean(r[4]),
                "used": True,
                "person": {"name": fix.get("deceased") or clean(r[5]),
                           "burial": r[3].date().isoformat()},
            }

    merged_person = merged_owner = 0
    for pid, rec in records.items():
        f = plots.get(pid)
        if f is None:
            sys.exit(f"ERROR: workbook space {pid} has no polygon on the map")
        p = f["properties"]
        expect = {"occupied", "buried"} if rec["used"] else {"vacant"}
        if p["status"] not in expect:
            sys.exit(f"ERROR: {pid} is {p['status']} on the map but "
                     f"{'used' if rec['used'] else 'not used'} in the workbook — "
                     "re-run digitize_layout.py export (check COLOR_OVERRIDES)")
        p.pop("owner", None)
        p.pop("person", None)
        if rec["used"]:
            p["person"] = rec["person"]
            merged_person += 1
        elif rec["owner"]:
            p["owner"] = rec["owner"]  # published: owner approved for vacant plots
            merged_owner += 1

    # -- 3. validate + write --
    vacant = [f for f in fc["features"] if f["properties"]["status"] == "vacant"]
    no_owner = [f["properties"]["id"] for f in vacant if not f["properties"].get("owner")]
    occupied = [f for f in fc["features"] if f["properties"]["status"] == "occupied"]
    no_person = [f["properties"]["id"] for f in occupied if not f["properties"].get("person")]
    fc["features"].sort(key=lambda f: (f["properties"]["lot"], f["properties"]["space"]))

    PLOTS.write_text(json.dumps(fc, indent=1))
    print(f"plots: {len(fc['features'])} | statuses changed: {changed or 'none'} | added: {added or 'none'}")
    print(f"merged: {merged_person} burials, {merged_owner} vacant-plot owners")
    print(f"vacant without owner: {no_owner or 'none'}")
    print(f"occupied without person: {no_person or 'none'}")


if __name__ == "__main__":
    main()
